"""Train, evaluate, save plots/model, and generate an A4 report."""

from __future__ import annotations

import argparse
from pathlib import Path

import joblib
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt

from modeling import TARGET_COLUMN, generate_demo_dataset, run_experiment

ROOT = Path(__file__).resolve().parent


def _add_chart(document: Document, path: Path, caption: str) -> None:
    document.add_picture(str(path), width=Inches(5.8))
    paragraph = document.add_paragraph(caption)
    paragraph.style = "Caption"


def create_report(result: dict, source_description: str, report_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    document.styles["Normal"].font.size = Pt(10)

    info = result["data_info"]
    metrics = result["metrics"]
    document.add_heading("Road accident severity prediction", 0)
    document.add_paragraph("Decision tree and model comparison project")
    document.add_heading("1. Problem statement", level=1)
    document.add_paragraph(
        "This project studies how road, vehicle, weather, and time conditions can be "
        "used to classify accident severity. The goal is an educational classification "
        "workflow, not an operational safety or emergency-response system."
    )
    document.add_heading("2. Dataset description", level=1)
    document.add_paragraph(
        f"Source: {source_description}. The input contains {info['original_rows']:,} rows "
        f"before cleaning and {info['rows_used']:,} rows after removing "
        f"{info['duplicate_rows']:,} duplicates and records without a target. "
        f"It has {len(result['feature_columns'])} features and "
        f"{len(info['class_counts'])} severity classes."
    )
    document.add_paragraph("Class counts: " + ", ".join(
        f"{label}: {count:,}" for label, count in info["class_counts"].items()
    ))
    document.add_heading("3. Preprocessing", level=1)
    document.add_paragraph(
        "Numeric features use median imputation followed by StandardScaler. Text "
        "features use most-frequent imputation and one-hot encoding. These transforms "
        "are inside each scikit-learn Pipeline and are fitted separately in each "
        "training fold, avoiding preprocessing leakage. The data is split 80/20 with "
        "stratification and random_state=42."
    )
    document.add_paragraph(
        f"Detected numeric features: {', '.join(info['numeric_columns']) or 'none'}. "
        f"Detected categorical features: {', '.join(info['categorical_columns']) or 'none'}."
    )
    document.add_page_break()

    document.add_heading("4. Algorithm selection and training", level=1)
    document.add_paragraph(
        "Logistic regression provides a linear baseline; a decision tree is included "
        "for interpretable rule-based classification; random forest provides a "
        "non-linear ensemble comparison. Each candidate is compared using stratified "
        "five-fold cross-validation on the training split only. Macro F1 is the "
        "selection metric so each class has equal influence. The test split is kept "
        "out of model selection."
    )
    for name, values in result["cv_results"].items():
        document.add_paragraph(
            f"{name}: mean CV macro F1 = {values['mean_f1_macro']:.3f} "
            f"(standard deviation {values['std_f1_macro']:.3f})"
        )
    document.add_paragraph(
        f"Selected model: {result['selected_model']} (highest mean training CV macro F1)."
    )
    document.add_heading("5. Held-out test results", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Test score"
    for metric_name, value in metrics.items():
        row = table.add_row().cells
        row[0].text = metric_name.replace("_", " ").title()
        row[1].text = f"{value:.3f}" if pd.notna(value) else "Not available"
    document.add_paragraph(
        "Precision reflects the reliability of predicted classes; recall reflects "
        "how many actual cases are found; macro F1 balances both while weighting "
        "each severity equally. ROC-AUC summarizes one-vs-rest ranking performance."
    )
    document.add_page_break()

    document.add_heading("6. Graphs and captions", level=1)
    charts_dir = report_path.parent / "plots"
    _add_chart(document, charts_dir / "class_distribution.png",
               "Figure 1. Number of records in each accident severity class.")
    _add_chart(document, charts_dir / "correlation_heatmap.png",
               "Figure 2. Pairwise correlations among numeric input features.")
    document.add_page_break()

    _add_chart(document, charts_dir / "confusion_matrix.png",
               "Figure 3. Held-out test predictions versus actual severity.")
    _add_chart(document, charts_dir / "roc_curve.png",
               "Figure 4. One-vs-rest ROC curves for the selected model.")
    document.add_page_break()

    _add_chart(document, charts_dir / "model_comparison.png",
               "Figure 5. Training-only cross-validation macro F1; bars show fold standard deviation.")
    document.add_heading("7. Analysis, limitations, and course outcome", level=1)
    document.add_paragraph(
        f"The cross-validation-selected model is {result['selected_model']}; the "
        "held-out scores above are reported only for final evaluation. Review the "
        "confusion matrix to identify which severity classes are confused. In a real "
        "deployment, missed high-severity cases may be more consequential than false "
        "alarms, so fatal/serious-class recall and the cost of errors deserve "
        "particular attention. This project does not establish that the model is "
        "safe or suitable for decisions about individuals or emergency response."
    )
    if source_description.startswith("synthetic"):
        document.add_paragraph(
            "Important limitation: this bundled demonstration dataset is synthetic "
            "and its labels are generated from a simple risk-score rule. Its metrics "
            "are not evidence of real-world predictive performance. Replace it with "
            "a documented, real accident dataset before drawing substantive conclusions."
        )
    document.add_paragraph(
        "CO5 – Solve real life problems using appropriate machine learning models and "
        "evaluate the performance measures. The project maps to CO5 by formulating "
        "accident severity as a classification problem, comparing appropriate "
        "models, selecting by cross-validation, and evaluating multiple measures."
    )
    document.save(report_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, help="CSV file to use instead of the demo dataset.")
    parser.add_argument("--target", default=TARGET_COLUMN, help="Name of the severity target column.")
    args = parser.parse_args()

    data_dir = ROOT / "data"
    output_dir = ROOT / "outputs"
    plots_dir = output_dir / "plots"
    data_dir.mkdir(exist_ok=True)
    output_dir.mkdir(exist_ok=True)
    if args.data:
        dataset = pd.read_csv(args.data)
        source_description = f"CSV supplied by the user ({args.data.name})"
    else:
        dataset = generate_demo_dataset()
        source_description = "synthetic demonstration data generated by this project"
        dataset.to_csv(data_dir / "road_accidents_demo.csv", index=False)

    result = run_experiment(dataset, args.target, plots_dir)
    artifact = {
        "model": result["model"],
        "target_column": args.target,
        "feature_columns": result["feature_columns"],
        "class_names": result["class_names"],
        "selected_model": result["selected_model"],
        "metrics": result["metrics"],
    }
    joblib.dump(artifact, output_dir / "road_accident_severity_model.joblib")
    create_report(result, source_description, output_dir / "road_accident_severity_report.docx")
    print(f"Rows used: {result['data_info']['rows_used']}")
    print(f"Selected by 5-fold CV macro F1: {result['selected_model']}")
    print(pd.DataFrame(result["cv_results"]).T.to_string())
    print(pd.Series(result["metrics"]).to_string())
    print(f"Plots: {plots_dir}")
    print(f"Model: {output_dir / 'road_accident_severity_model.joblib'}")
    print(f"Report: {output_dir / 'road_accident_severity_report.docx'}")


if __name__ == "__main__":
    main()
